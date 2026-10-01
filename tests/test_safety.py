"""Offline regressions for failures found in the independent review."""

import unittest
from unittest.mock import patch

from common.matcher import locality_score
from common import lookup_log
from v2 import geopy_distance as geo


class SafetyTests(unittest.TestCase):
    def test_city_code_and_joined_name(self):
        self.assertEqual(locality_score("ANNA NAGAR", "K-4 ANNANAGAR"), 100)
        self.assertLess(locality_score("NAGAR", "K-4 ANNANAGAR"), 86)

    def test_stack_detector_keeps_distinct_rows(self):
        point = (10.0, 78.0)
        coords = {
            ("DISTRICT A", "TOWN PS"): point,
            ("DISTRICT A", "RURAL PS"): point,
            ("DISTRICT B", "UNRELATED"): (11.0, 79.0),
        }
        self.assertEqual(geo.coordinate_collisions(coords),
                         {("DISTRICT A", "TOWN PS"),
                          ("DISTRICT A", "RURAL PS")})

    def test_address_geocode_keeps_neighbouring_pin_result(self):
        # A real lookup (PIN 502032) was matched by Google to an exact street
        # in the neighbouring PIN 502033. Discarding it left the officer with
        # no answer at all, so the address geocode must accept it.
        neighbouring_pin = {
            "types": ["street_address", "subpremise"],
            "geometry": {"location": {"lat": 17.5, "lng": 78.3}},
            "formatted_address": "Some street, Telangana 502033, India",
            "address_components": [
                {"long_name": "502033", "types": ["postal_code"]},
                {"long_name": "Telangana", "types": ["administrative_area_level_1"]},
            ],
        }
        with patch.object(geo, "_geocode_top", return_value=neighbouring_pin):
            self.assertEqual(geo.geocode_address("Some street, 502032"),
                             ((17.5, 78.3), "Some street, Telangana 502033, India"))
            self.assertEqual(geo.geocode_address("Some street, 502032",
                                                 return_state=True)[2], "Telangana")

    def test_distance_lookup_does_not_retry_blank_station(self):
        cache = {("DISTRICT", "MISSING"): None}
        with patch.object(geo, "_load_coords_cache", return_value=cache), \
             patch.object(geo, "geocode_address", return_value=((13.0, 80.0), "address")), \
             patch.object(geo, "geocode_station") as paid:
            result = geo.rank_ps_by_distance("address", ["MISSING"], "DISTRICT")
        self.assertEqual(result, [])
        self.assertEqual(result.excluded, ["MISSING"])
        paid.assert_not_called()

    def test_distance_lookup_uses_stateful_station_coordinates(self):
        station = ("SHARED DISTRICT", "CENTRAL")
        cache = {station: (13.0, 80.0)}
        stateful = {
            ("TELANGANA", *station): (17.0, 78.0),
            ("ANDHRA PRADESH", *station): (13.0, 80.0),
        }
        with patch.object(geo, "_load_coords_cache", return_value=cache), \
             patch.object(geo, "_stateful_coords_cache", stateful), \
             patch.object(geo, "geocode_address",
                          return_value=((17.0, 78.0), "address", "Telangana")):
            result = geo.measure_from_address(
                "address", [("TELANGANA", *station),
                            ("ANDHRA PRADESH", *station)])
        self.assertEqual(result[0]["distance_km"], 0)
        self.assertGreater(result[1]["distance_km"], 0)
        self.assertEqual(result.address_state, "Telangana")

    def test_bulk_write_pauses_new_different_name_collision(self):
        existing = {("DISTRICT", "TOWN"): (13.0, 80.0)}
        with patch.object(geo, "_load_coords_cache", return_value=existing), \
             patch.object(geo.openpyxl, "load_workbook") as load_workbook:
            geo._write_coords([("DISTRICT", "RURAL", 13.0, 80.0)])
        load_workbook.assert_not_called()

    def test_unrouted_lookup_is_included_in_measurement_log(self):
        with patch.object(lookup_log, "_append_row") as append:
            lookup_log.log_lookup({"case": 0, "results": []},
                                  "unmatched address", "", "", interactive=False)
        values = append.call_args.args[1]
        self.assertEqual(values["RESULT LOOKUP"], "NO RESULT")
        self.assertEqual(values["RESULT MATCH"], "No suggestion")
        self.assertEqual(values["PREDICTED PS"], "")
        self.assertEqual(values["SHOWN STATIONS"], "")

    def test_log_records_every_shown_station(self):
        results = [
            {"rank": 1, "police_station": "BORABANDA", "district": "HYDERABAD",
             "distance": "~1.2 km", "confidence": "HIGH"},
            {"rank": 2, "police_station": "SR NAGAR", "district": "HYDERABAD",
             "distance": "~2.9 km", "confidence": "LOW",
             "coordinate_unverified": True},
            {"rank": 3, "police_station": "KUKATPALLY", "district": "CYBERABAD",
             "distance": "N/A", "confidence": "LOW"},
        ]
        self.assertEqual(
            lookup_log.shown_summary(results),
            "1. BORABANDA (HYDERABAD) ~1.2 km | "
            "2. SR NAGAR (HYDERABAD) ~2.9 km [shared point] | "
            "3. KUKATPALLY (CYBERABAD)")

    def test_shown_column_added_without_moving_existing_data(self):
        import os
        import tempfile
        import openpyxl
        from openpyxl.worksheet.table import Table

        headers = ["FILE NO", "ADDRESS", "ACTUAL PS KNOWN", "PREDICTED PS",
                   "PV STATUS", "MATCH", "RESULT LOOKUP", "RESULT MATCH",
                   "PREDICTED DISTRICT"]
        old_row = ["F1", "old address", "TOWN", "TOWN", "DONE", "YES",
                   "DISTRICT", "Very Likely", "DIST"]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LookupLogs"
        ws.append(headers)
        ws.append(old_row)
        ws.add_table(Table(displayName="Table2", ref="A1:I2"))
        path = os.path.join(tempfile.mkdtemp(), "log.xlsx")
        wb.save(path)

        values = {"ADDRESS": "new address", "PREDICTED PS": "A",
                  "PREDICTED DISTRICT": "D", "RESULT LOOKUP": "DISTRICT",
                  "RESULT MATCH": "Likely", "SHOWN STATIONS": "1. A (D) ~1 km"}
        lookup_log._append_row(path, values)
        lookup_log._append_row(path, values)   # header must not be added twice

        ws = openpyxl.load_workbook(path)["LookupLogs"]
        self.assertEqual([c.value for c in ws[1]], headers + ["SHOWN STATIONS"])
        self.assertEqual([c.value for c in ws[2]][:9], old_row)
        self.assertIsNone(ws.cell(row=2, column=10).value)
        self.assertEqual(ws.cell(row=3, column=10).value, "1. A (D) ~1 km")
        self.assertEqual(ws.cell(row=4, column=2).value, "new address")
        self.assertIsNone(ws.cell(row=3, column=1).value)   # FILE NO stays manual


if __name__ == "__main__":
    unittest.main()
